from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from personal_qa.match_analysis import MatchAnalysisService, safe_rate  # noqa: E402

MATCH_FIELDS = (
    "match_id",
    "game_start_utc",
    "game_duration_seconds",
    "queue_id",
    "map_id",
    "game_version",
    "is_remake",
)
PARTICIPANT_FIELDS = (
    "match_id",
    "player_id",
    "participant_id",
    "team_id",
    "champion_id",
    "champion_name",
    "team_position",
    "win",
    "kills",
    "deaths",
    "assists",
    "kda",
    "cs",
    "cs_per_min",
    "gold_earned",
    "gold_per_min",
    "damage_to_champions",
    "damage_per_min",
    "damage_taken",
    "damage_taken_per_min",
    "vision_score",
    "vision_per_min",
    "objective_damage",
    "objective_damage_per_min",
    "is_target_player",
    "game_start_utc",
    "unavailable_reasons",
)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _safe_members(archive: zipfile.ZipFile) -> list[zipfile.ZipInfo]:
    members = archive.infolist()
    for member in members:
        path = PurePosixPath(member.filename.replace("\\", "/"))
        if path.is_absolute() or ".." in path.parts or (path.parts and ":" in path.parts[0]):
            raise ValueError(f"unsafe ZIP member: {member.filename}")
    return members


def _one_member(
    members: list[zipfile.ZipInfo], *, basename: str | None = None, suffix: str | None = None
) -> zipfile.ZipInfo:
    found = []
    for member in members:
        name = PurePosixPath(member.filename.replace("\\", "/")).name
        if basename == name or (suffix and name.casefold().endswith(suffix.casefold())):
            found.append(member)
    if len(found) != 1:
        label = basename or suffix
        raise ValueError(f"expected exactly one {label!r} member, found {len(found)}")
    return found[0]


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV encoding is not supported")


def _bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip().casefold()
    if text in {"true", "1", "yes"}:
        return True
    if text in {"false", "0", "no"}:
        return False
    return None


def _number(value: Any, *, integer: bool = False) -> int | float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if integer else number


def _utc_from_timestamp(value: Any) -> str | None:
    number = _number(value)
    if number is None:
        return None
    if number > 10_000_000_000:
        number /= 1000
    return datetime.fromtimestamp(number, UTC).isoformat().replace("+00:00", "Z")


def _jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n" for row in rows),
        encoding="utf-8",
    )


def _target_participant(
    source: dict[str, Any], alias: str, game_start_utc: str | None
) -> dict[str, Any]:
    duration = _number(source.get("game_duration"), integer=True)
    kills = _number(source.get("kills"), integer=True)
    deaths = _number(source.get("deaths"), integer=True)
    assists = _number(source.get("assists"), integer=True)
    cs = _number(source.get("total_cs"), integer=True)
    gold = _number(source.get("gold_earned"), integer=True)
    damage = _number(source.get("total_damage_dealt_to_champions"), integer=True)
    damage_taken = _number(source.get("total_damage_taken"), integer=True)
    vision = _number(source.get("vision_score"), integer=True)
    objective = _number(source.get("damage_dealt_to_objectives"), integer=True)
    return {
        "match_id": alias,
        "player_id": "TARGET_PLAYER",
        "participant_id": _number(source.get("participant_id"), integer=True),
        "team_id": _number(source.get("team_id"), integer=True),
        "champion_id": _number(source.get("champion_id"), integer=True),
        "champion_name": source.get("champion_name"),
        "team_position": source.get("team_position") or None,
        "win": _bool(source.get("win")),
        "kills": kills,
        "deaths": deaths,
        "assists": assists,
        "kda": None
        if kills is None or deaths is None or assists is None
        else round((kills + assists) / max(1, deaths), 4),
        "cs": cs,
        "cs_per_min": safe_rate(cs, duration),
        "gold_earned": gold,
        "gold_per_min": safe_rate(gold, duration),
        "damage_to_champions": damage,
        "damage_per_min": safe_rate(damage, duration),
        "damage_taken": damage_taken,
        "damage_taken_per_min": safe_rate(damage_taken, duration),
        "vision_score": vision,
        "vision_per_min": safe_rate(vision, duration),
        "objective_damage": objective,
        "objective_damage_per_min": safe_rate(objective, duration),
        "is_target_player": True,
        "game_start_utc": game_start_utc,
        "unavailable_reasons": {},
    }


def _xlsx_participant(
    source: dict[str, Any], alias: str, player_id: str, game_start_utc: str | None
) -> dict[str, Any]:
    duration = _number(source.get("game_duration"), integer=True)
    kills = _number(source.get("kills"), integer=True)
    deaths = _number(source.get("deaths"), integer=True)
    assists = _number(source.get("assists"), integer=True)
    gold = _number(source.get("gold_earned"), integer=True)
    damage = _number(source.get("total_damage_dealt_to_champions"), integer=True)
    damage_taken = _number(source.get("total_damage_taken"), integer=True)
    vision = _number(source.get("vision_score"), integer=True)
    objective = _number(source.get("damage_dealt_to_objectives"), integer=True)
    return {
        "match_id": alias,
        "player_id": player_id,
        "participant_id": _number(source.get("participant_id"), integer=True),
        "team_id": _number(source.get("team_id"), integer=True),
        "champion_id": _number(source.get("champion_id"), integer=True),
        "champion_name": source.get("champion_name"),
        "team_position": source.get("team_position") or None,
        "win": _bool(source.get("win")),
        "kills": kills,
        "deaths": deaths,
        "assists": assists,
        "kda": None
        if kills is None or deaths is None or assists is None
        else round((kills + assists) / max(1, deaths), 4),
        "cs": None,
        "cs_per_min": None,
        "gold_earned": gold,
        "gold_per_min": safe_rate(gold, duration),
        "damage_to_champions": damage,
        "damage_per_min": safe_rate(damage, duration),
        "damage_taken": damage_taken,
        "damage_taken_per_min": safe_rate(damage_taken, duration),
        "vision_score": vision,
        "vision_per_min": safe_rate(vision, duration),
        "objective_damage": objective,
        "objective_damage_per_min": safe_rate(objective, duration),
        "is_target_player": False,
        "game_start_utc": game_start_utc,
        "unavailable_reasons": {
            "cs": "not_present_in_cleaned_workbook",
            "cs_per_min": "not_present_in_cleaned_workbook",
        },
    }


def prepare_personal_data(
    outer_zip: Path, output_dir: Path, local_mapping_path: Path
) -> dict[str, Any]:
    source_before = _sha256(outer_zip.read_bytes())
    with zipfile.ZipFile(outer_zip) as archive:
        members = _safe_members(archive)
        if archive.testzip() is not None:
            raise ValueError("outer ZIP failed CRC validation")
        matches_member = _one_member(members, basename="matches.csv")
        workbook_member = _one_member(members, suffix=".xlsx")
        matches_rows = list(
            csv.DictReader(io.StringIO(_decode_csv(archive.read(matches_member)), newline=""))
        )
        workbook_bytes = archive.read(workbook_member)

    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    if "league_data_cleaned" not in workbook.sheetnames:
        raise ValueError("league_data_cleaned sheet is missing")
    sheet = workbook["league_data_cleaned"]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(iterator)]
    cleaned_rows = [
        dict(zip(headers, row, strict=False))
        for row in iterator
        if any(value is not None for value in row)
    ]
    workbook.close()

    if len(matches_rows) != 30 or len(headers) != 36 or len(cleaned_rows) != 180:
        raise ValueError("unexpected personal source dimensions")
    raw_ids = [row.get("") or row.get("match_id") or row.get("game_id") for row in matches_rows]
    if any(not isinstance(value, str) or not re.fullmatch(r"KR_\d+", value) for value in raw_ids):
        raise ValueError("matches.csv does not contain the expected Match-V5 IDs")
    match_aliases = {raw: f"MATCH_{index:03d}" for index, raw in enumerate(raw_ids, start=1)}
    target_by_match = {raw: row for raw, row in zip(raw_ids, matches_rows, strict=True)}
    cleaned_by_match: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in cleaned_rows:
        cleaned_by_match[str(row["game_id"])].append(row)
    overlap = set(raw_ids) & set(cleaned_by_match)
    if len(overlap) != 18 or any(len(cleaned_by_match[raw]) != 10 for raw in overlap):
        raise ValueError("the expected 18-match derived relationship was not found")

    target_source_ids = {str(row.get("puuid")) for row in matches_rows if row.get("puuid")}
    for raw in overlap:
        target_pid = int(float(target_by_match[raw]["participant_id"]))
        target_source_ids.update(
            str(row["puuid"])
            for row in cleaned_by_match[raw]
            if int(row["participant_id"]) == target_pid
        )
    player_aliases: dict[str, str] = {
        value: "TARGET_PLAYER" for value in sorted(target_source_ids)
    }
    next_player = 1
    for row in cleaned_rows:
        source_id = str(row.get("puuid"))
        if source_id and source_id not in player_aliases:
            player_aliases[source_id] = f"PLAYER_{next_player:05d}"
            next_player += 1

    matches: list[dict[str, Any]] = []
    participants: list[dict[str, Any]] = []
    for raw, source in zip(raw_ids, matches_rows, strict=True):
        alias = match_aliases[raw]
        start = _utc_from_timestamp(source.get("game_start_timestamp"))
        duration = _number(source.get("game_duration"), integer=True)
        matches.append(
            {
                "match_id": alias,
                "game_start_utc": start,
                "game_duration_seconds": duration,
                "queue_id": _number(source.get("queue_id"), integer=True),
                "map_id": _number(source.get("map_id"), integer=True),
                "game_version": source.get("game_version"),
                "is_remake": bool(
                    duration is not None
                    and duration < 300
                    or _bool(source.get("game_ended_in_early_surrender"))
                ),
            }
        )
        target_pid = int(float(source["participant_id"]))
        if raw in overlap:
            for cleaned in cleaned_by_match[raw]:
                participant_id = int(cleaned["participant_id"])
                if participant_id == target_pid:
                    participants.append(_target_participant(source, alias, start))
                else:
                    source_player = str(cleaned.get("puuid"))
                    participants.append(
                        _xlsx_participant(
                            cleaned, alias, player_aliases[source_player], start
                        )
                    )
        else:
            participants.append(_target_participant(source, alias, start))

    service = MatchAnalysisService(matches, participants)
    features = []
    opponent_statuses = Counter()
    for match in matches:
        match_id = match["match_id"]
        target = next(
            row
            for row in participants
            if row["match_id"] == match_id and row["player_id"] == "TARGET_PLAYER"
        )
        comparison = service.compare_with_position_opponent(match_id, "TARGET_PLAYER")
        opponent_statuses[comparison["opponent_match_status"]] += 1
        deltas = {
            item["metric"]: item["delta"] for item in comparison.get("metric_comparisons", [])
        }
        opponent = comparison.get("opponent") or {}
        features.append(
            {
                "match_id": match_id,
                "player_id": "TARGET_PLAYER",
                "champion_name": target.get("champion_name"),
                "team_position": target.get("team_position"),
                "opponent_player_id": opponent.get("player_id"),
                "opponent_champion_name": opponent.get("champion_name"),
                "kda_delta": deltas.get("kda"),
                "cs_per_min_delta": deltas.get("cs_per_min"),
                "gold_per_min_delta": deltas.get("gold_per_min"),
                "damage_per_min_delta": deltas.get("damage_per_min"),
                "vision_per_min_delta": deltas.get("vision_per_min"),
                "objective_damage_per_min_delta": deltas.get("objective_damage_per_min"),
                "opponent_match_status": comparison["opponent_match_status"],
                "opponent_match_reason": comparison.get("reason"),
            }
        )

    full_match_counts = Counter(row["match_id"] for row in participants)
    target_counts = Counter(
        row["match_id"] for row in participants if row["player_id"] == "TARGET_PLAYER"
    )
    winners = Counter(row["match_id"] for row in participants if row.get("win") is True)
    full_ids = {match_aliases[raw] for raw in overlap}
    if any(full_match_counts[match_id] != 10 for match_id in full_ids):
        raise ValueError("full participant match does not contain 10 rows")
    if any(winners[match_id] != 5 for match_id in full_ids):
        raise ValueError("full participant match does not contain five winners")
    if any(target_counts[match["match_id"]] != 1 for match in matches):
        raise ValueError("target player is not unique")
    if any(tuple(row) != MATCH_FIELDS for row in matches):
        raise ValueError("matches field order mismatch")
    if any(tuple(row) != PARTICIPANT_FIELDS for row in participants):
        raise ValueError("participant field order mismatch")

    output_dir.mkdir(parents=True, exist_ok=True)
    local_mapping_path.parent.mkdir(parents=True, exist_ok=True)
    _jsonl(output_dir / "matches.jsonl", matches)
    _jsonl(output_dir / "match_participants.jsonl", participants)
    _jsonl(output_dir / "player_match_features.jsonl", features)
    mapping = {
        "warning": "local-only sensitive mapping; never commit or share",
        "matches": match_aliases,
        "players": player_aliases,
    }
    local_mapping_path.write_text(
        json.dumps(mapping, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    generated_text = "\n".join(
        path.read_text(encoding="utf-8")
        for path in output_dir.glob("*.jsonl")
    )
    leaked = [value for value in [*match_aliases, *player_aliases] if value in generated_text]
    report = {
        "valid": not leaked,
        "source_archive_sha256": source_before,
        "source_dimensions": {
            "matches_csv": {"rows": len(matches_rows), "columns": len(matches_rows[0])},
            "cleaned_workbook": {"rows": len(cleaned_rows), "columns": len(headers)},
            "overlapping_derived_matches": len(overlap),
        },
        "outputs": {
            "matches": len(matches),
            "match_participants": len(participants),
            "full_participant_matches": len(full_ids),
            "target_only_matches": len(matches) - len(full_ids),
            "player_match_features": len(features),
        },
        "opponent_match_status": dict(sorted(opponent_statuses.items())),
        "target_player_exactly_once": True,
        "full_matches_have_10_participants": True,
        "full_matches_have_5_winners": True,
        "rates_recomputed": True,
        "zero_duration_matches": sum(
            match["game_duration_seconds"] in (None, 0) for match in matches
        ),
        "short_or_remake_matches": sum(match["is_remake"] for match in matches),
        "zero_duration_safe": all(
            row["game_duration_seconds"] not in (None, 0) for row in matches
        )
        or all(
            row["cs_per_min"] is None
            for row in participants
            if next(
                match["game_duration_seconds"]
                for match in matches
                if match["match_id"] == row["match_id"]
            )
            in (None, 0)
        ),
        "sensitive_identifier_findings": len(leaked),
        "actual_match_ids_in_outputs": len(re.findall(r"\bKR_\d+\b", generated_text)),
        "source_archive_unchanged": _sha256(outer_zip.read_bytes()) == source_before,
        "limitations": [
            "Only 18 of 30 target matches have ten-participant rows in the cleaned workbook.",
            "Opponent CS is unavailable because the cleaned workbook has no CS columns.",
            "The large IRON timeline file was inspected but not transformed.",
        ],
    }
    if not report["valid"] or report["actual_match_ids_in_outputs"]:
        raise ValueError("sensitive identifiers remain in generated outputs")
    (output_dir / "validation_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare and validate local personal QA tables.")
    parser.add_argument("--outer-zip", type=Path, required=True)
    parser.add_argument(
        "--output-dir", type=Path, default=PROJECT_ROOT / "data" / "processed" / "personal_qa"
    )
    parser.add_argument(
        "--local-mapping",
        type=Path,
        default=PROJECT_ROOT / "data" / "local" / "anonymization_map.local.json",
    )
    args = parser.parse_args()
    report = prepare_personal_data(
        args.outer_zip.resolve(), args.output_dir.resolve(), args.local_mapping.resolve()
    )
    print(
        json.dumps(
            {
                "status": "ok",
                "outputs": report["outputs"],
                "opponent_match_status": report["opponent_match_status"],
                "sensitive_identifier_findings": report["sensitive_identifier_findings"],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
